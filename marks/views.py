from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib import messages
from django.db.models import Sum, Avg, Count, Q, F, Max, Min
from django.http import JsonResponse, HttpResponse
from django.core.paginator import Paginator
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.urls import reverse_lazy
from django.utils import timezone
from django.db import transaction, IntegrityError
from decimal import Decimal, InvalidOperation
import csv
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from io import BytesIO
from datetime import date

from .models import (
    AcademicYear, Class, Subject, Teacher, Student,
    Exam, Marks, Result, Attendance, Achievement, StudyMaterial,
    Assignment, AssignmentSubmission, DisciplinaryAction, PerformanceShare,
    StudentRole, StudentLeave, StudentClassChangeRequest
)


def home(request):
    if request.user.is_authenticated:
        return redirect('marks:dashboard')
    return redirect('marks:login')


def _teacher_for(request):
    return getattr(request.user, 'teacher_profile', None) or getattr(request.user, 'teacher', None)


def _student_for(request):
    return getattr(request.user, 'student_profile', None)


def _parent_for(request):
    return getattr(request.user, 'parent_profile', None)


def _is_staff(user):
    return user.is_authenticated and user.is_staff


def _teacher_can_access_class(teacher, student_class):
    return (
        student_class.class_teacher_id == teacher.id
        or teacher.assigned_classes.filter(id=student_class.id).exists()
    )


def _teacher_subjects_for_class(teacher, student_class):
    return student_class.subjects.filter(teachers=teacher).distinct()


class AdminRequiredMixin:
    def dispatch(self, request, *args, **kwargs):
        if not _is_staff(request.user):
            return redirect('marks:portal_dashboard')
        return super().dispatch(request, *args, **kwargs)


@login_required
def portal_dashboard(request):
    teacher = _teacher_for(request)
    student = _student_for(request)
    parent = _parent_for(request)

    if teacher and (not teacher.is_active or teacher.approval_status != 'approved'):
        return render(request, 'marks/portal_dashboard.html', {
            'role': 'pending', 'approval_status': teacher.get_approval_status_display(),
        })

    if teacher:
        classes = Class.objects.filter(Q(class_teacher=teacher) | Q(subject_teachers=teacher)).distinct().prefetch_related('students')
        assignments = Assignment.objects.filter(teacher=teacher).select_related('student_class', 'subject')[:6]
        materials = StudyMaterial.objects.filter(uploaded_by=teacher)[:6]
        mark_entry_options = [
            {'exam': exam, 'student_class': student_class}
            for exam in Exam.objects.filter(classes__in=classes, is_published=False).distinct()[:12]
            for student_class in exam.classes.filter(id__in=classes.values('id'))
        ]
        return render(request, 'marks/portal_dashboard.html', {
            'role': 'teacher', 'teacher': teacher, 'classes': classes,
            'assignments': assignments, 'materials': materials,
            'student_count': Student.objects.filter(student_class__class_teacher=teacher, is_active=True).count(),
            'pending_submissions': AssignmentSubmission.objects.filter(assignment__teacher=teacher, status='submitted').count(),
            'teacher_roles': [teacher.get_role_display()] + ([teacher.duties] if teacher.duties else []),
            'mark_entry_options': mark_entry_options,
        })

    if student:
        results = Result.objects.filter(student=student, exam__is_published=True).select_related('exam')
        attendance = student.attendance_records.all()
        assignments = Assignment.objects.filter(student_class=student.student_class, is_published=True).select_related('subject')[:8]
        submissions = {submission.assignment_id: submission for submission in student.assignment_submissions.all()}
        materials = StudyMaterial.objects.filter(student_class=student.student_class, is_published=True).select_related('subject')[:8]
        return render(request, 'marks/portal_dashboard.html', {
            'role': 'student', 'student': student, 'results': results,
            'attendance': attendance, 'attendance_total': attendance.count(),
            'attendance_present': attendance.filter(status__in=['present', 'late']).count(),
            'assignments': assignments, 'submissions': submissions, 'materials': materials,
            'achievements': student.achievements.all()[:6],
            'discipline': student.disciplinary_actions.filter(parent_visible=True)[:6],
            'student_roles': student.roles.filter(is_active=True),
            'leaves': student.leaves.filter(status='approved')[:6],
        })

    if parent:
        students = parent.students.filter(is_active=True).select_related('student_class')
        selected_student = get_object_or_404(students, id=request.GET.get('student')) if request.GET.get('student') else students.first()
        context = {'role': 'parent', 'parent': parent, 'students': students, 'student': selected_student}
        if selected_student:
            context.update({
                'results': Result.objects.filter(student=selected_student, exam__is_published=True).select_related('exam'),
                'attendance': selected_student.attendance_records.all()[:30],
                'attendance_total': selected_student.attendance_records.count(),
                'attendance_present': selected_student.attendance_records.filter(status__in=['present', 'late']).count(),
                'achievements': selected_student.achievements.all()[:6],
                'discipline': selected_student.disciplinary_actions.filter(parent_visible=True)[:6],
                'student_roles': selected_student.roles.filter(is_active=True),
                'leaves': selected_student.leaves.filter(status='approved')[:6],
                'materials': StudyMaterial.objects.filter(student_class=selected_student.student_class, is_published=True).select_related('subject')[:8],
            })
        return render(request, 'marks/portal_dashboard.html', context)

    return redirect('marks:dashboard')


@login_required
def teacher_attendance(request, class_id):
    teacher = _teacher_for(request)
    if not teacher and not request.user.is_staff:
        return redirect('marks:portal_dashboard')
    student_class = get_object_or_404(Class, id=class_id, class_teacher=teacher)
    attendance_date = request.POST.get('date') or request.GET.get('date') or date.today().isoformat()
    students = student_class.students.filter(is_active=True)
    if request.method == 'POST':
        for student in students:
            Attendance.objects.update_or_create(
                student=student,
                date=attendance_date,
                defaults={
                    'status': request.POST.get(f'status_{student.id}', 'present'),
                    'remarks': request.POST.get(f'remarks_{student.id}', ''),
                    'marked_by': teacher,
                },
            )
        messages.success(request, 'Attendance saved for the class.')
        return redirect('marks:teacher_attendance', class_id=class_id)
    records = {record.student_id: record for record in Attendance.objects.filter(student__in=students, date=attendance_date)}
    return render(request, 'marks/teacher_attendance.html', {'student_class': student_class, 'students': students, 'records': records, 'attendance_date': attendance_date})


@login_required
def teacher_create_achievement(request):
    teacher = _teacher_for(request)
    if not teacher and not request.user.is_staff:
        return redirect('marks:portal_dashboard')
    if request.method == 'POST':
        Achievement.objects.create(
            student_id=request.POST['student'], title=request.POST['title'], category=request.POST['category'],
            description=request.POST.get('description', ''), achieved_on=request.POST['achieved_on'],
            issuer=request.POST.get('issuer', ''), evidence_url=request.POST.get('evidence_url', ''), created_by=teacher,
        )
        messages.success(request, 'Achievement added to the student record.')
        return redirect('marks:portal_dashboard')
    return render(request, 'marks/teacher_record_form.html', {'record_type': 'Achievement', 'students': Student.objects.filter(is_active=True), 'categories': Achievement.CATEGORY_CHOICES})


@login_required
def teacher_create_discipline(request):
    teacher = _teacher_for(request)
    if not teacher and not request.user.is_staff:
        return redirect('marks:portal_dashboard')
    if request.method == 'POST':
        DisciplinaryAction.objects.create(
            student_id=request.POST['student'], action_type=request.POST['action_type'], severity=request.POST['severity'],
            incident_date=request.POST['incident_date'], description=request.POST['description'],
            resolution=request.POST.get('resolution', ''), parent_visible=request.POST.get('parent_visible') == 'on', reported_by=teacher,
        )
        messages.success(request, 'Student wellbeing record saved.')
        return redirect('marks:portal_dashboard')
    return render(request, 'marks/teacher_record_form.html', {'record_type': 'Disciplinary record', 'students': Student.objects.filter(is_active=True), 'action_types': DisciplinaryAction.ACTION_CHOICES, 'severity_levels': DisciplinaryAction.SEVERITY_CHOICES})


@login_required
def teacher_create_material(request):
    teacher = _teacher_for(request)
    if not teacher and not request.user.is_staff:
        return redirect('marks:portal_dashboard')
    if request.method == 'POST':
        StudyMaterial.objects.create(
            title=request.POST['title'], description=request.POST.get('description', ''), subject_id=request.POST.get('subject') or None,
            student_class_id=request.POST.get('student_class') or None, external_url=request.POST.get('external_url', ''),
            file=request.FILES.get('file'), is_published=request.POST.get('is_published') == 'on', uploaded_by=teacher,
        )
        messages.success(request, 'Study material published.')
        return redirect('marks:portal_dashboard')
    return render(request, 'marks/teacher_record_form.html', {'record_type': 'Study material', 'subjects': Subject.objects.all(), 'classes': Class.objects.all()})


@login_required
def teacher_create_assignment(request):
    teacher = _teacher_for(request)
    if not teacher and not request.user.is_staff:
        return redirect('marks:portal_dashboard')
    if request.method == 'POST':
        Assignment.objects.create(
            title=request.POST['title'], instructions=request.POST['instructions'],
            subject_id=request.POST.get('subject') or None, student_class_id=request.POST['student_class'],
            due_date=request.POST['due_date'], attachment=request.FILES.get('attachment'),
            is_published=request.POST.get('is_published') == 'on', teacher=teacher,
        )
        messages.success(request, 'Assignment published.')
        return redirect('marks:portal_dashboard')
    classes = Class.objects.filter(class_teacher=teacher) if teacher else Class.objects.all()
    return render(request, 'marks/teacher_record_form.html', {'record_type': 'Assignment', 'subjects': Subject.objects.all(), 'classes': classes})


@login_required
def student_submit_assignment(request, assignment_id):
    student = _student_for(request)
    assignment = get_object_or_404(Assignment, id=assignment_id, student_class=student.student_class, is_published=True)
    if request.method == 'POST':
        AssignmentSubmission.objects.update_or_create(
            assignment=assignment, student=student,
            defaults={'response': request.POST.get('response', ''), 'attachment': request.FILES.get('attachment'), 'status': 'submitted'},
        )
        messages.success(request, 'Assignment submitted.')
    return redirect('marks:portal_dashboard')


@login_required
def create_performance_share(request, student_id):
    student = get_object_or_404(Student, id=student_id)
    if not (_teacher_for(request) or request.user.is_staff or _parent_for(request)):
        return redirect('marks:portal_dashboard')
    share = PerformanceShare.objects.create(student=student, created_by=request.user)
    return render(request, 'marks/share_created.html', {'share': share, 'student': student})


def shared_performance(request, token):
    share = get_object_or_404(PerformanceShare.objects.select_related('student'), token=token, is_active=True)
    if share.expires_at and share.expires_at <= timezone.now():
        return HttpResponse('This performance link has expired.', status=410)
    student = share.student
    return render(request, 'marks/shared_performance_branded.html', {
        'student': student,
        'results': Result.objects.filter(student=student, exam__is_published=True).select_related('exam'),
        'attendance': student.attendance_records.all()[:30],
        'achievements': student.achievements.all()[:8],
    })


@login_required
def dashboard(request):
    teacher = _teacher_for(request)
    if teacher and (not teacher.is_active or teacher.approval_status != 'approved'):
        return render(request, 'marks/portal_dashboard.html', {'role': 'pending', 'approval_status': teacher.get_approval_status_display()})

    if teacher and not request.user.is_staff:
        classes = Class.objects.filter(Q(class_teacher=teacher) | Q(subject_teachers=teacher)).distinct()
    else:
        classes = Class.objects.all()
    selected_class = classes.filter(id=request.GET.get('class')).first() if request.GET.get('class') else classes.first()
    exams = Exam.objects.filter(classes=selected_class).order_by('-start_date') if selected_class else Exam.objects.none()
    selected_exam = exams.filter(id=request.GET.get('exam')).first() if request.GET.get('exam') else exams.first()
    if request.method == 'POST':
        if not request.user.is_staff:
            return redirect('marks:dashboard')
        exam_name = request.POST.get('name', '').strip()
        exam_date = request.POST.get('exam_date')
        class_id = request.POST.get('class_id')
        if exam_name and exam_date and class_id:
            student_class = get_object_or_404(Class, id=class_id)
            academic_year = student_class.academic_year
            exam = Exam.objects.create(
                name=exam_name, exam_type='unit_test', academic_year=academic_year,
                start_date=exam_date, end_date=exam_date,
            )
            exam.classes.add(student_class)
            messages.success(request, f'{exam.name} created for {student_class}.')
            return redirect(f'{reverse_lazy("marks:dashboard")}?class={student_class.id}&exam={exam.id}')
    active_year = AcademicYear.objects.filter(is_active=True).first()
    return render(request, 'marks/marks_workspace.html', {
        'classes': classes, 'selected_class': selected_class, 'exams': exams,
        'selected_exam': selected_exam, 'active_year': active_year,
        'is_admin': request.user.is_staff,
    })


@user_passes_test(_is_staff)
def school_setup(request):
    if request.method == 'POST':
        action = request.POST.get('action')
        try:
            if action == 'add_class':
                Class.objects.create(
                    name=request.POST['name'].strip(), section=request.POST['section'],
                    academic_year_id=request.POST['academic_year'],
                )
                messages.success(request, 'Class and section added.')
            elif action == 'edit_class':
                class_obj = get_object_or_404(Class, id=request.POST['class_id'])
                class_obj.name = request.POST['name'].strip()
                class_obj.section = request.POST['section']
                class_obj.academic_year_id = request.POST['academic_year']
                class_obj.save()
                messages.success(request, 'Class updated.')
            elif action == 'add_subject':
                subject = Subject.objects.create(
                    name=request.POST['name'].strip(), code=request.POST['code'].strip(),
                    max_marks=request.POST['max_marks'], pass_marks=request.POST['pass_marks'],
                )
                subject.classes.add(request.POST['class_id'])
                if request.POST.get('teacher_id'):
                    teacher = get_object_or_404(Teacher, id=request.POST['teacher_id'])
                    teacher.subjects.add(subject)
                    teacher.assigned_classes.add(request.POST['class_id'])
                messages.success(request, 'Subject mark template added.')
            elif action == 'edit_subject':
                subject = get_object_or_404(Subject, id=request.POST['subject_id'])
                subject.name = request.POST['name'].strip()
                subject.code = request.POST['code'].strip()
                subject.max_marks = request.POST['max_marks']
                subject.pass_marks = request.POST['pass_marks']
                subject.save()
                subject.classes.set([request.POST['class_id']])
                if request.POST.get('teacher_id'):
                    teacher = get_object_or_404(Teacher, id=request.POST['teacher_id'])
                    teacher.subjects.add(subject)
                    teacher.assigned_classes.add(request.POST['class_id'])
                messages.success(request, 'Subject mark template updated.')
            elif action == 'add_student':
                Student.objects.create(
                    admission_number=request.POST['admission_number'].strip(),
                    roll_number=request.POST['roll_number'].strip(),
                    first_name=request.POST['first_name'].strip(), last_name=request.POST['last_name'].strip(),
                    gender=request.POST['gender'], date_of_birth=request.POST['date_of_birth'],
                    student_class_id=request.POST['student_class'], father_name=request.POST['father_name'].strip(),
                    mother_name=request.POST['mother_name'].strip(), address=request.POST['address'].strip(),
                    phone=request.POST['phone'].strip(), email=request.POST.get('email', '').strip(),
                    admission_date=request.POST['admission_date'], approval_status='approved', is_active=True,
                )
                messages.success(request, 'Student added.')
            elif action == 'edit_student':
                student = get_object_or_404(Student, id=request.POST['student_id'])
                student.first_name = request.POST['first_name'].strip()
                student.last_name = request.POST['last_name'].strip()
                student.roll_number = request.POST['roll_number'].strip()
                student.student_class_id = request.POST['student_class']
                student.phone = request.POST['phone'].strip()
                student.email = request.POST.get('email', '').strip()
                student.save(update_fields=['first_name', 'last_name', 'roll_number', 'student_class', 'phone', 'email'])
                messages.success(request, 'Student updated.')
        except (KeyError, ValueError, IntegrityError) as error:
            messages.error(request, f'Could not save that change: {error}')
        return redirect('marks:school_setup')

    return render(request, 'marks/school_setup.html', {
        'years': AcademicYear.objects.all(), 'classes': Class.objects.select_related('academic_year'),
        'subjects': Subject.objects.prefetch_related('classes', 'teachers'),
        'students': Student.objects.select_related('student_class')[:100],
        'teachers': Teacher.objects.filter(is_active=True, approval_status='approved'),
    })


class StudentRegistrationView(CreateView):
    model = Student
    template_name = 'marks/student_form.html'
    fields = ['admission_number', 'roll_number', 'first_name', 'last_name', 'gender',
              'date_of_birth', 'student_class', 'father_name', 'mother_name',
              'address', 'phone', 'email', 'admission_date', 'photo']
    success_url = reverse_lazy('marks:home')

    def form_valid(self, form):
        form.instance.approval_status = 'pending'
        form.instance.is_active = False
        messages.success(self.request, 'Student registration submitted for admin approval.')
        return super().form_valid(form)


class TeacherRegistrationView(CreateView):
    model = Teacher
    template_name = 'marks/teacher_form.html'
    fields = ['employee_id', 'first_name', 'last_name', 'email', 'phone', 'role', 'duties']
    success_url = reverse_lazy('marks:home')

    def form_valid(self, form):
        form.instance.approval_status = 'pending'
        form.instance.is_active = False
        messages.success(self.request, 'Teacher registration submitted for admin approval.')
        return super().form_valid(form)


@login_required
def request_class_change(request):
    student = _student_for(request)
    if not student:
        return redirect('marks:portal_dashboard')
    if request.method == 'POST':
        StudentClassChangeRequest.objects.create(
            student=student,
            requested_class_id=request.POST['requested_class'],
            reason=request.POST.get('reason', ''),
        )
        messages.success(request, 'Class change request submitted for admin approval.')
        return redirect('marks:portal_dashboard')
    return render(request, 'marks/class_change_request_form.html', {'classes': Class.objects.exclude(id=student.student_class_id)})


@user_passes_test(_is_staff)
def approve_class_change(request, request_id):
    change_request = get_object_or_404(StudentClassChangeRequest, id=request_id, status='pending')
    if request.method == 'POST':
        decision = request.POST.get('decision')
        change_request.status = decision if decision in ('approved', 'rejected') else 'rejected'
        change_request.reviewed_by = request.user
        change_request.reviewed_at = timezone.now()
        if change_request.status == 'approved':
            change_request.student.student_class = change_request.requested_class
            change_request.student.save(update_fields=['student_class'])
        change_request.save(update_fields=['status', 'reviewed_by', 'reviewed_at'])
        messages.success(request, f'Class change request {change_request.status}.')
    return redirect('marks:dashboard')


@user_passes_test(_is_staff)
def approve_student(request, student_id):
    student = get_object_or_404(Student, id=student_id, approval_status='pending')
    if request.method == 'POST':
        decision = request.POST.get('decision')
        student.approval_status = decision if decision in ('approved', 'rejected') else 'rejected'
        student.is_active = student.approval_status == 'approved'
        student.save(update_fields=['approval_status', 'is_active'])
        messages.success(request, f'Student registration {student.approval_status}.')
    return redirect('marks:dashboard')


@user_passes_test(_is_staff)
def approve_teacher(request, teacher_id):
    teacher = get_object_or_404(Teacher, id=teacher_id, approval_status='pending')
    if request.method == 'POST':
        decision = request.POST.get('decision')
        teacher.approval_status = decision if decision in ('approved', 'rejected') else 'rejected'
        teacher.is_active = teacher.approval_status == 'approved'
        teacher.save(update_fields=['approval_status', 'is_active'])
        messages.success(request, f'Teacher registration {teacher.approval_status}.')
    return redirect('marks:dashboard')


class AcademicYearListView(AdminRequiredMixin, ListView):
    model = AcademicYear
    template_name = 'marks/academic_year_list.html'
    context_object_name = 'years'
    paginate_by = 10


class AcademicYearCreateView(AdminRequiredMixin, CreateView):
    model = AcademicYear
    template_name = 'marks/academic_year_form.html'
    fields = ['name', 'start_date', 'end_date', 'is_active']
    success_url = reverse_lazy('marks:academic_year_list')


class AcademicYearUpdateView(AdminRequiredMixin, UpdateView):
    model = AcademicYear
    template_name = 'marks/academic_year_form.html'
    fields = ['name', 'start_date', 'end_date', 'is_active']
    success_url = reverse_lazy('marks:academic_year_list')


class AcademicYearDeleteView(AdminRequiredMixin, DeleteView):
    model = AcademicYear
    template_name = 'marks/academic_year_confirm_delete.html'
    success_url = reverse_lazy('marks:academic_year_list')


class ClassListView(AdminRequiredMixin, ListView):
    model = Class
    template_name = 'marks/class_list.html'
    context_object_name = 'classes'
    paginate_by = 20

    def get_queryset(self):
        queryset = Class.objects.select_related('academic_year', 'class_teacher').all()
        academic_year = self.request.GET.get('academic_year')
        if academic_year:
            queryset = queryset.filter(academic_year_id=academic_year)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['academic_years'] = AcademicYear.objects.all()
        return context


class ClassCreateView(AdminRequiredMixin, CreateView):
    model = Class
    template_name = 'marks/class_form.html'
    fields = ['name', 'section', 'academic_year', 'class_teacher']
    success_url = reverse_lazy('marks:class_list')


class ClassUpdateView(AdminRequiredMixin, UpdateView):
    model = Class
    template_name = 'marks/class_form.html'
    fields = ['name', 'section', 'academic_year', 'class_teacher']
    success_url = reverse_lazy('marks:class_list')


class ClassDeleteView(AdminRequiredMixin, DeleteView):
    model = Class
    template_name = 'marks/class_confirm_delete.html'
    success_url = reverse_lazy('marks:class_list')


class ClassDetailView(AdminRequiredMixin, DetailView):
    model = Class
    template_name = 'marks/class_detail.html'
    context_object_name = 'class_obj'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        students = self.object.students.filter(is_active=True)
        context['students'] = students
        context['subjects'] = self.object.subjects.all()
        return context


class SubjectListView(AdminRequiredMixin, ListView):
    model = Subject
    template_name = 'marks/subject_list.html'
    context_object_name = 'subjects'
    paginate_by = 20


class SubjectCreateView(AdminRequiredMixin, CreateView):
    model = Subject
    template_name = 'marks/subject_form.html'
    fields = ['name', 'code', 'subject_type', 'max_marks', 'pass_marks', 'classes']
    success_url = reverse_lazy('marks:subject_list')


class SubjectUpdateView(AdminRequiredMixin, UpdateView):
    model = Subject
    template_name = 'marks/subject_form.html'
    fields = ['name', 'code', 'subject_type', 'max_marks', 'pass_marks', 'classes']
    success_url = reverse_lazy('marks:subject_list')


class SubjectDeleteView(AdminRequiredMixin, DeleteView):
    model = Subject
    template_name = 'marks/subject_confirm_delete.html'
    success_url = reverse_lazy('marks:subject_list')


class TeacherListView(AdminRequiredMixin, ListView):
    model = Teacher
    template_name = 'marks/teacher_list.html'
    context_object_name = 'teachers'
    paginate_by = 20


class TeacherCreateView(AdminRequiredMixin, CreateView):
    model = Teacher
    template_name = 'marks/teacher_form.html'
    fields = ['employee_id', 'first_name', 'last_name', 'email', 'phone', 'subjects', 'assigned_classes', 'role', 'duties', 'is_active', 'approval_status']
    success_url = reverse_lazy('marks:teacher_list')


class TeacherUpdateView(AdminRequiredMixin, UpdateView):
    model = Teacher
    template_name = 'marks/teacher_form.html'
    fields = ['employee_id', 'first_name', 'last_name', 'email', 'phone', 'subjects', 'assigned_classes', 'role', 'duties', 'is_active', 'approval_status']
    success_url = reverse_lazy('marks:teacher_list')


class TeacherDeleteView(AdminRequiredMixin, DeleteView):
    model = Teacher
    template_name = 'marks/teacher_confirm_delete.html'
    success_url = reverse_lazy('marks:teacher_list')


class StudentListView(AdminRequiredMixin, ListView):
    model = Student
    template_name = 'marks/student_list.html'
    context_object_name = 'students'
    paginate_by = 30

    def get_queryset(self):
        queryset = Student.objects.select_related('student_class', 'student_class__academic_year').filter(is_active=True)
        student_class = self.request.GET.get('class')
        if student_class:
            queryset = queryset.filter(student_class_id=student_class)
        search = self.request.GET.get('search')
        if search:
            queryset = queryset.filter(
                Q(first_name__icontains=search) |
                Q(last_name__icontains=search) |
                Q(admission_number__icontains=search) |
                Q(roll_number__icontains=search)
            )
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['classes'] = Class.objects.all()
        return context


class StudentCreateView(AdminRequiredMixin, CreateView):
    model = Student
    template_name = 'marks/student_form.html'
    fields = ['admission_number', 'roll_number', 'first_name', 'last_name', 'gender',
              'date_of_birth', 'student_class', 'father_name', 'mother_name',
              'address', 'phone', 'email', 'admission_date', 'photo', 'approval_status']
    success_url = reverse_lazy('marks:student_list')


class StudentUpdateView(AdminRequiredMixin, UpdateView):
    model = Student
    template_name = 'marks/student_form.html'
    fields = ['admission_number', 'roll_number', 'first_name', 'last_name', 'gender',
              'date_of_birth', 'student_class', 'father_name', 'mother_name',
              'address', 'phone', 'email', 'admission_date', 'photo', 'is_active', 'approval_status']
    success_url = reverse_lazy('marks:student_list')


class StudentDeleteView(AdminRequiredMixin, DeleteView):
    model = Student
    template_name = 'marks/student_confirm_delete.html'
    success_url = reverse_lazy('marks:student_list')


class StudentDetailView(AdminRequiredMixin, DetailView):
    model = Student
    template_name = 'marks/student_detail.html'
    context_object_name = 'student'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['marks'] = Marks.objects.filter(student=self.object).select_related('subject', 'exam', 'exam__academic_year')
        context['results'] = Result.objects.filter(student=self.object).select_related('exam')
        return context


class ExamListView(AdminRequiredMixin, ListView):
    model = Exam
    template_name = 'marks/exam_list.html'
    context_object_name = 'exams'
    paginate_by = 20

    def get_queryset(self):
        queryset = Exam.objects.prefetch_related('classes').all()
        academic_year = self.request.GET.get('academic_year')
        if academic_year:
            queryset = queryset.filter(academic_year_id=academic_year)
        return queryset

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['academic_years'] = AcademicYear.objects.all()
        return context


class ExamCreateView(AdminRequiredMixin, CreateView):
    model = Exam
    template_name = 'marks/exam_form.html'
    fields = ['name', 'exam_type', 'academic_year', 'classes', 'start_date', 'end_date', 'weightage']
    success_url = reverse_lazy('marks:exam_list')


class ExamUpdateView(AdminRequiredMixin, UpdateView):
    model = Exam
    template_name = 'marks/exam_form.html'
    fields = ['name', 'exam_type', 'academic_year', 'classes', 'start_date', 'end_date', 'weightage', 'is_published']
    success_url = reverse_lazy('marks:exam_list')


class ExamDeleteView(AdminRequiredMixin, DeleteView):
    model = Exam
    template_name = 'marks/exam_confirm_delete.html'
    success_url = reverse_lazy('marks:exam_list')


class ExamDetailView(AdminRequiredMixin, DetailView):
    model = Exam
    template_name = 'marks/exam_detail.html'
    context_object_name = 'exam'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['classes'] = self.object.classes.all()
        return context


@login_required
def marks_entry(request, exam_id, class_id):
    teacher = _teacher_for(request)
    if not request.user.is_staff and (not teacher or not teacher.is_active or teacher.approval_status != 'approved'):
        return redirect('marks:dashboard')
    exam = get_object_or_404(Exam, id=exam_id)
    student_class = get_object_or_404(Class, id=class_id, exams=exam)
    students = student_class.students.filter(is_active=True)
    subjects = student_class.subjects.all()
    if teacher and not request.user.is_staff:
        if not _teacher_can_access_class(teacher, student_class):
            return redirect('marks:dashboard')
        subjects = _teacher_subjects_for_class(teacher, student_class)

    if request.method == 'POST':
        with transaction.atomic():
            for student in students:
                for subject in subjects:
                    marks_key = f'marks_{student.id}_{subject.id}'
                    absent_key = f'absent_{student.id}_{subject.id}'
                    remarks_key = f'remarks_{student.id}_{subject.id}'

                    marks_value = request.POST.get(marks_key)
                    is_absent = request.POST.get(absent_key) == 'on'
                    remarks = request.POST.get(remarks_key, '')

                    if marks_value or is_absent:
                        try:
                            numeric_value = Decimal(marks_value or '0')
                        except InvalidOperation:
                            messages.error(request, f'Invalid mark entered for {student.full_name}, {subject.name}.')
                            return redirect('marks:marks_entry', exam_id=exam.id, class_id=student_class.id)
                        if numeric_value < 0 or numeric_value > subject.max_marks:
                            messages.error(request, f'Marks for {subject.name} must be between 0 and {subject.max_marks}.')
                            return redirect('marks:marks_entry', exam_id=exam.id, class_id=student_class.id)
                        Marks.objects.update_or_create(
                            student=student,
                            subject=subject,
                            exam=exam,
                            defaults={
                                'marks_obtained': numeric_value if not is_absent else 0,
                                'is_absent': is_absent,
                                'graded_by': teacher,
                                'remarks': remarks,
                            }
                        )
        messages.success(request, 'Marks saved successfully!')
        return redirect('marks:marks_entry', exam_id=exam.id, class_id=student_class.id)

    existing_marks = Marks.objects.filter(exam=exam, student__in=students).select_related('student', 'subject')
    marks_dict = {(m.student_id, m.subject_id): m for m in existing_marks}
    student_summaries = {}
    for student in students:
        student_marks = [marks_dict.get((student.id, subject.id)) for subject in subjects]
        student_marks = [mark for mark in student_marks if mark]
        total = sum((mark.marks_obtained for mark in student_marks if not mark.is_absent), Decimal('0'))
        max_total = sum(subject.max_marks for subject in subjects)
        student_summaries[student.id] = {
            'total': total,
            'max_total': max_total,
            'percentage': round((float(total) / max_total) * 100, 2) if max_total else 0,
            'complete': len(student_marks) == subjects.count(),
        }
    all_marks_complete = all(
        summary['complete'] for summary in student_summaries.values()
    ) if students.exists() and subjects.exists() else False
    full_subjects = student_class.subjects.all()
    all_class_marks_complete = all(
        Marks.objects.filter(exam=exam, student=student, subject__in=full_subjects).count() == full_subjects.count()
        for student in students
    ) if students.exists() and full_subjects.exists() else False

    context = {
        'exam': exam,
        'student_class': student_class,
        'students': students,
        'subjects': subjects,
        'marks_dict': marks_dict,
        'student_summaries': student_summaries,
        'all_marks_complete': all_marks_complete,
        'can_generate_report': (request.user.is_staff or bool(teacher)) and all_class_marks_complete,
    }
    return render(request, 'marks/marks_entry.html', context)


@login_required
def calculate_results(request, exam_id):
    teacher = _teacher_for(request)
    if not request.user.is_staff and (not teacher or not teacher.is_active or teacher.approval_status != 'approved'):
        return redirect('marks:dashboard')
    exam = get_object_or_404(Exam, id=exam_id)
    target_classes = exam.classes.all()
    if not request.user.is_staff:
        target_classes = target_classes.filter(Q(class_teacher=teacher) | Q(subject_teachers=teacher)).distinct()

    if request.method == 'POST':
        missing = []
        for student_class in target_classes:
            subjects = student_class.subjects.all()
            for student in student_class.students.filter(is_active=True):
                existing_subject_ids = set(Marks.objects.filter(student=student, exam=exam).values_list('subject_id', flat=True))
                missing.extend(subject.name for subject in subjects if subject.id not in existing_subject_ids)
        if missing:
            messages.error(request, f'Cannot publish yet. {len(missing)} student subject marks are still missing.')
            return redirect('marks:calculate_results', exam_id=exam.id)
        with transaction.atomic():
            for student_class in target_classes:
                students = student_class.students.filter(is_active=True)
                for student in students:
                    marks = Marks.objects.filter(student=student, exam=exam)
                    if marks.exists():
                        result, created = Result.objects.update_or_create(
                            student=student,
                            exam=exam,
                            defaults={}
                        )
                        result.calculate_result()

        if request.user.is_staff:
            exam.is_published = True
            exam.save()
            messages.success(request, 'Results calculated and published successfully!')
        else:
            messages.success(request, 'Class reports generated successfully.')
        return redirect('marks:dashboard')

    context = {'exam': exam}
    return render(request, 'marks/calculate_results.html', context)


def _can_view_student(request, student):
    if request.user.is_staff:
        return True
    own_student = _student_for(request)
    if own_student and own_student.id == student.id:
        return True
    parent = _parent_for(request)
    if parent and parent.students.filter(id=student.id).exists():
        return True
    teacher = _teacher_for(request)
    return bool(teacher and _teacher_can_access_class(teacher, student.student_class))


@login_required
def result_sheet(request, exam_id, class_id):
    exam = get_object_or_404(Exam, id=exam_id)
    student_class = get_object_or_404(Class, id=class_id)
    students = student_class.students.filter(is_active=True)

    results = Result.objects.filter(exam=exam, student__in=students).select_related('student')
    results = sorted(results, key=lambda x: float(x.percentage), reverse=True)

    for idx, result in enumerate(results, 1):
        result.rank = idx
        result.save()

    subjects = student_class.subjects.all()

    context = {
        'exam': exam,
        'student_class': student_class,
        'results': results,
        'subjects': subjects,
    }
    return render(request, 'marks/result_sheet.html', context)


@login_required
def student_report_card(request, student_id, exam_id):
    student = get_object_or_404(Student, id=student_id)
    if not _can_view_student(request, student):
        return redirect('marks:dashboard')
    exam = get_object_or_404(Exam, id=exam_id)
    marks = Marks.objects.filter(student=student, exam=exam).select_related('subject')
    result = Result.objects.filter(student=student, exam=exam).first()

    if not result:
        result = Result(student=student, exam=exam)
        result.calculate_result()

    context = {
        'student': student,
        'exam': exam,
        'marks': marks,
        'result': result,
        'school_name': 'Jawahar Navodaya Vidyalaya, Ernakulam',
    }
    return render(request, 'marks/report_card.html', context)


@login_required
def export_marks_excel(request, exam_id, class_id):
    exam = get_object_or_404(Exam, id=exam_id)
    student_class = get_object_or_404(Class, id=class_id)
    teacher = _teacher_for(request)
    if not request.user.is_staff and (not teacher or not _teacher_can_access_class(teacher, student_class)):
        return redirect('marks:dashboard')
    students = student_class.students.filter(is_active=True)
    subjects = student_class.subjects.all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{exam.name} Marks"

    header_font = Font(bold=True, size=12)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font_white = Font(bold=True, size=12, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    ws.merge_cells('A1:F1')
    ws['A1'] = 'JAWAHAR NAVODAYA VIDYALAYA, ERNAKULAM'
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A2:F2')
    ws['A2'] = f'{exam.name} - {exam.exam_type}'
    ws['A2'].font = Font(bold=True, size=14)
    ws['A2'].alignment = Alignment(horizontal='center')

    ws.merge_cells('A3:F3')
    ws['A3'] = f'Class: {student_class.name}{student_class.section} | Academic Year: {student_class.academic_year}'
    ws['A3'].font = Font(size=12)
    ws['A3'].alignment = Alignment(horizontal='center')

    row = 5
    headers = ['Roll No', 'Admission No', 'Student Name']
    for subject in subjects:
        headers.append(f'{subject.name}\n({subject.max_marks})')
    headers.extend(['Total', 'Max Total', 'Percentage', 'Grade', 'Rank'])

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font_white
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    result_by_student = {result.student_id: result for result in Result.objects.filter(exam=exam, student__in=students)}
    student_rows = []
    for student in students:
        marks_dict = {mark.subject_id: mark for mark in Marks.objects.filter(student=student, exam=exam)}
        total = sum((mark.marks_obtained for mark in marks_dict.values() if not mark.is_absent), Decimal('0'))
        max_total = sum(subject.max_marks for subject in subjects)
        percentage = round((float(total) / max_total) * 100, 2) if max_total else 0
        student_rows.append((student, marks_dict, result_by_student.get(student.id), total, max_total, percentage))
    student_rows.sort(key=lambda row: row[5], reverse=True)

    for idx, (student, marks_dict, result, total, max_total, percentage) in enumerate(student_rows, 1):
        row += 1

        ws.cell(row=row, column=1, value=student.roll_number).border = thin_border
        ws.cell(row=row, column=1).alignment = center_align
        ws.cell(row=row, column=2, value=student.admission_number).border = thin_border
        ws.cell(row=row, column=2).alignment = center_align
        ws.cell(row=row, column=3, value=student.full_name).border = thin_border

        col = 4
        for subject in subjects:
            mark = marks_dict.get(subject.id)
            if mark:
                if mark.is_absent:
                    cell_value = 'AB'
                else:
                    cell_value = f"{mark.marks_obtained}/{subject.max_marks}"
            else:
                cell_value = '-'
            cell = ws.cell(row=row, column=col, value=cell_value)
            cell.border = thin_border
            cell.alignment = center_align
            col += 1

        ws.cell(row=row, column=col, value=float(result.total_marks) if result else float(total)).border = thin_border
        ws.cell(row=row, column=col).alignment = center_align
        col += 1
        ws.cell(row=row, column=col, value=result.max_total_marks if result else max_total).border = thin_border
        ws.cell(row=row, column=col).alignment = center_align
        col += 1
        ws.cell(row=row, column=col, value=f"{result.percentage if result else percentage}%").border = thin_border
        ws.cell(row=row, column=col).alignment = center_align
        col += 1
        ws.cell(row=row, column=col, value=result.grade if result else '-').border = thin_border
        ws.cell(row=row, column=col).alignment = center_align
        col += 1
        ws.cell(row=row, column=col, value=result.rank if result else idx).border = thin_border
        ws.cell(row=row, column=col).alignment = center_align

    for column_index, column_cells in enumerate(ws.columns, 1):
        max_length = 0
        column = openpyxl.utils.get_column_letter(column_index)
        for cell in column_cells:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 20)
        ws.column_dimensions[column].width = adjusted_width

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{exam.name}_{student_class.name}{student_class.section}_marks.xlsx"'
    wb.save(response)
    return response


@login_required
def export_report_card_pdf(request, student_id, exam_id):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    student = get_object_or_404(Student, id=student_id)
    if not _can_view_student(request, student):
        return redirect('marks:dashboard')
    exam = get_object_or_404(Exam, id=exam_id)
    marks = Marks.objects.filter(student=student, exam=exam).select_related('subject')
    result = Result.objects.filter(student=student, exam=exam).first()
    if not result:
        result = Result(student=student, exam=exam)
        result.calculate_result()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="{student.admission_number}_{exam.name}_report_card.pdf"'
    )

    doc = SimpleDocTemplate(
        response,
        pagesize=A4,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'SchoolTitle', parent=styles['Title'], fontSize=16, spaceAfter=2,
    )
    sub_style = ParagraphStyle(
        'SchoolSub', parent=styles['Normal'], alignment=1,
        textColor=colors.HexColor('#475569'), spaceAfter=10,
    )

    story = [
        Paragraph('Jawahar Navodaya Vidyalaya, Ernakulam', title_style),
        Paragraph(
            'PM SHRI School &middot; Neriamangalam, Ernakulam Dt., Kerala &middot; CBSE Affiliated',
            sub_style,
        ),
        Paragraph(
            f'Report Card &mdash; {exam.name} ({exam.get_exam_type_display()})',
            styles['Heading2'],
        ),
        Spacer(1, 6),
        Paragraph(
            f'<b>Name:</b> {student.full_name} &nbsp;&nbsp; '
            f'<b>Admission No:</b> {student.admission_number} &nbsp;&nbsp; '
            f'<b>Class:</b> {student.student_class} &nbsp;&nbsp; '
            f'<b>Roll No:</b> {student.roll_number}',
            styles['Normal'],
        ),
        Spacer(1, 14),
    ]

    table_data = [['Subject', 'Max Marks', 'Marks Obtained', 'Grade']]
    for mark in marks:
        table_data.append([
            mark.subject.name,
            str(mark.subject.max_marks),
            'AB' if mark.is_absent else str(mark.marks_obtained),
            '-' if mark.is_absent else mark.grade,
        ])
    table_data.append(['Total', str(result.max_total_marks), str(result.total_marks), result.grade])

    table = Table(table_data, colWidths=[70 * mm, 35 * mm, 40 * mm, 25 * mm])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f1f5f9')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#f8fafc')]),
    ]))
    story.append(table)
    story.append(Spacer(1, 14))
    story.append(Paragraph(
        f'<b>Percentage:</b> {result.percentage}% &nbsp;&nbsp; '
        f'<b>Result:</b> {"PASS" if result.is_pass else "FAIL"} &nbsp;&nbsp; '
        f'<b>Rank:</b> {result.rank or "-"}',
        styles['Normal'],
    ))

    doc.build(story)
    return response


@login_required
def class_performance(request, class_id):
    student_class = get_object_or_404(Class, id=class_id)
    exams = Exam.objects.filter(classes=student_class).order_by('-start_date')

    exam_data = []
    for exam in exams:
        students = student_class.students.filter(is_active=True)
        results = Result.objects.filter(exam=exam, student__in=students)
        if results.exists():
            avg_percentage = results.aggregate(avg=Avg('percentage'))['avg']
            pass_count = results.filter(is_pass=True).count()
            total_count = results.count()
            pass_percentage = (pass_count / total_count * 100) if total_count > 0 else 0
            exam_data.append({
                'exam': exam,
                'avg_percentage': round(avg_percentage, 2) if avg_percentage else 0,
                'pass_percentage': round(pass_percentage, 2),
                'total_students': total_count,
            })

    context = {
        'student_class': student_class,
        'exam_data': exam_data,
    }
    return render(request, 'marks/class_performance.html', context)


@login_required
def subject_analysis(request, subject_id):
    subject = get_object_or_404(Subject, id=subject_id)
    exams = Exam.objects.filter(classes__subjects=subject).distinct().order_by('-start_date')

    exam_analysis = []
    for exam in exams:
        marks = Marks.objects.filter(subject=subject, exam=exam).select_related('student')
        if marks.exists():
            total_students = marks.count()
            present_students = marks.filter(is_absent=False).count()
            absent_students = total_students - present_students
            passed = marks.filter(is_absent=False, marks_obtained__gte=subject.pass_marks).count()
            failed = present_students - passed
            avg_marks = marks.filter(is_absent=False).aggregate(avg=Avg('marks_obtained'))['avg']
            max_marks = marks.filter(is_absent=False).aggregate(max=Max('marks_obtained'))['max']
            min_marks = marks.filter(is_absent=False).aggregate(min=Min('marks_obtained'))['min']

            exam_analysis.append({
                'exam': exam,
                'total_students': total_students,
                'present_students': present_students,
                'absent_students': absent_students,
                'passed': passed,
                'failed': failed,
                'pass_percentage': round((passed / present_students * 100), 2) if present_students > 0 else 0,
                'avg_marks': round(avg_marks, 2) if avg_marks else 0,
                'max_marks': max_marks,
                'min_marks': min_marks,
            })

    context = {
        'subject': subject,
        'exam_analysis': exam_analysis,
    }
    return render(request, 'marks/subject_analysis.html', context)


@login_required
def api_student_marks(request, student_id, exam_id):
    student = get_object_or_404(Student, id=student_id)
    exam = get_object_or_404(Exam, id=exam_id)
    marks = Marks.objects.filter(student=student, exam=exam).select_related('subject')

    data = []
    for mark in marks:
        data.append({
            'subject': mark.subject.name,
            'subject_code': mark.subject.code,
            'max_marks': mark.subject.max_marks,
            'marks_obtained': float(mark.marks_obtained),
            'percentage': mark.percentage,
            'grade': mark.grade,
            'is_absent': mark.is_absent,
        })

    result = Result.objects.filter(student=student, exam=exam).first()
    if result:
        data.append({
            'total_marks': float(result.total_marks),
            'max_total_marks': result.max_total_marks,
            'percentage': float(result.percentage),
            'grade': result.grade,
            'rank': result.rank,
            'is_pass': result.is_pass,
        })

    return JsonResponse({'marks': data})